"""Erzeugt die Protokollbögen für den Klassenrat als HTML (Druckvorlage) und Word.

Beide Ausgaben entstehen aus derselben Beschreibung (siehe `grundschule` und
`mittelschule` unten), damit Papier- und Word-Fassung immer gleich aufgebaut sind.

Aufruf:  python3 klassenrat/werkzeuge/bogen_erzeugen.py
Danach:  node klassenrat/werkzeuge/pdf_erzeugen.mjs   (HTML -> PDF)
"""

from html import escape
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent

import sys
sys.path.insert(0, str(Path(__file__).resolve().parent))
import bilder  # noqa: E402

ICONS = bilder.ALLE
GESICHTER = [("froh", "gut"), ("mittel", "geht so"), ("traurig", "nicht gut")]
PNG = ROOT / "bilder" / "png"
EIGENES_LOGO = ROOT / "bilder" / "logo.png"  # optional: Originallogo der Schule

# Akzentfarben je Schulart (aus dem Logo)
FARBEN = {
    "gs": {"akzent": bilder.ROT, "hell": "#FBEAEC", "zweit": bilder.GELB},
    "ms": {"akzent": bilder.GRUEN, "hell": "#E3F3EE", "zweit": bilder.GRAU},
}


# ---------------------------------------------------------------------------
# HTML-Ausgabe
# ---------------------------------------------------------------------------

class HtmlBogen:
    def __init__(self, titel, untertitel, stufe):
        self.titel, self.untertitel, self.stufe = titel, untertitel, stufe
        self.seiten = [[]]

    def _add(self, s):
        self.seiten[-1].append(s)

    @staticmethod
    def _icon(name):
        return f'<span class="ic">{ICONS[name]}</span>' if name else ""

    def kopf(self, felder):
        if EIGENES_LOGO.exists():
            logo = f'<img src="../bilder/logo.png" alt="Schullogo">'
        else:
            logo = ICONS["logo"]
        self._add(
            f'<header class="kopf"><div class="logo">{logo}</div>'
            f'<div class="titel"><h1>{escape(self.titel)}</h1><p>{escape(self.untertitel)}</p></div>'
            f'<div class="bild">{ICONS["klassenrat"]}</div></header>{self._felder_html(felder)}'
        )

    def ablauf(self, schritte):
        items = "".join(f"<li><b>{i}</b> {escape(s)}</li>" for i, s in enumerate(schritte, 1))
        self._add(f'<ol class="ablauf"><span>Ablauf:</span>{items}</ol>')

    def abschnitt(self, nr, titel, hinweis=None, icon=None):
        num = f'<span class="nr">{nr}</span>' if nr else ""
        tip = f'<span class="hinweis">{escape(hinweis)}</span>' if hinweis else ""
        self._add(f'<h2>{num}{self._icon(icon)}<span>{escape(titel)}</span>{tip}</h2>')

    def unterabschnitt(self, titel, hinweis=None, icon=None):
        tip = f'<span class="hinweis">{escape(hinweis)}</span>' if hinweis else ""
        self._add(f'<h3>{self._icon(icon)}<span>{escape(titel)}</span>{tip}</h3>')

    def block_anfang(self, titel, icon=None):
        self._add(f'<section class="block"><div class="blocktitel"><span class="ic-rund">{self._icon(icon)}</span>'
                  f'{escape(titel)}<span class="nrfeld">Nr. ______</span></div>')

    def block_ende(self):
        self._add("</section>")

    def _felder_html(self, felder):
        cells = "".join(
            f'<div class="feld" style="flex:{w}">{self._icon(ic)}<label>{escape(l)}</label><span class="wert"></span></div>'
            for l, w, *rest in felder
            for ic in [rest[0] if rest else None]
        )
        return f'<div class="felder">{cells}</div>'

    def felder(self, felder):
        self._add(self._felder_html(felder))

    def tabelle(self, spalten, zeilen):
        head = "".join(
            f'<th style="width:{w}%">{self._spaltenkopf(l)}</th>' for l, w, _ in self._prozent(spalten)
        )
        row = "".join(f"<td>{self._zelle(art)}</td>" for _, _, art in spalten)
        body = "".join(f"<tr>{row}</tr>" for _ in range(zeilen))
        self._add(f'<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>')

    def _spaltenkopf(self, label):
        if label.startswith("@"):  # Symbol-Spalte, z. B. "@daumen_hoch|dafür"
            ic, text = label[1:].split("|")
            return f'{self._icon(ic)}<span>{escape(text)}</span>'
        return escape(label)

    def _zelle(self, art):
        if art.startswith("wahl:"):
            return "".join(f'<span class="box"></span>{escape(o)} ' for o in art[5:].split("/"))
        if art == "gesichter":
            return "".join(f'<span class="gesicht">{self._icon(i)}</span>' for i, _ in GESICHTER)
        if art == "kasten":
            return '<span class="box"></span>'
        return ""

    @staticmethod
    def _prozent(spalten):
        total = sum(w for _, w, _ in spalten)
        return [(l, round(100 * w / total, 2), a) for l, w, a in spalten]

    def linien(self, n, label=None):
        lab = f"<label>{escape(label)}</label>" if label else ""
        self._add(f'<div class="linien">{lab}' + '<span class="linie"></span>' * n + "</div>")

    def ankreuzen(self, label, optionen):
        opts = "".join(f'<span class="opt"><span class="box"></span>{escape(o)}</span>' for o in optionen)
        lab = f"<label>{escape(label)}</label>" if label else ""
        self._add(f'<div class="ankreuzen">{lab}{opts}</div>')

    def gesichter_strichliste(self, label):
        opts = "".join(
            f'<div class="strich"><span class="gesicht gross">{self._icon(i)}</span>'
            f'<span class="striche"><small>{escape(t)}</small></span></div>'
            for i, t in GESICHTER
        )
        lab = f"<label>{escape(label)}</label>" if label else ""
        self._add(f'<div class="strichliste">{lab}{opts}</div>')

    def unterschriften(self, rollen):
        cells = "".join(f'<div><span></span><label>{escape(r)}</label></div>' for r in rollen)
        self._add(f'<div class="unterschriften">{cells}</div>')

    def hinweis(self, text):
        self._add(f'<p class="fuss-hinweis">{escape(text)}</p>')

    def seitenumbruch(self):
        self.seiten.append([])

    def speichern(self, pfad):
        n = len(self.seiten)
        seiten = "".join(
            f'<div class="seite">{"".join(inhalt)}'
            f'<footer>{escape(self.titel)} · {escape(self.untertitel)}<span>Seite {i} von {n}</span></footer></div>'
            for i, inhalt in enumerate(self.seiten, 1)
        )
        css = (ROOT / "werkzeuge" / "bogen.css").read_text(encoding="utf-8")
        f = FARBEN[self.stufe]
        farben = f"--akzent:{f['akzent']};--akzent-hell:{f['hell']};--zweit:{f['zweit']}"
        pfad.write_text(
            f'<!doctype html><html lang="de"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width, initial-scale=1">'
            f'<title>{escape(self.titel)} – {escape(self.untertitel)}</title>'
            f"<style>{css}</style></head><body class=\"{self.stufe}\" style=\"{farben}\">{seiten}</body></html>",
            encoding="utf-8",
        )


# ---------------------------------------------------------------------------
# Word-Ausgabe
# ---------------------------------------------------------------------------



def _rand(cell, **kanten):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for kante in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{kante}")
        wert = kanten.get(kante)
        if wert:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(wert))
            el.set(qn("w:color"), "555555")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)


def _schattierung(cell, farbe):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), farbe)
    cell._tc.get_or_add_tcPr().append(shd)


def _absatz_abstand(p, vor=0, nach=0):
    p.paragraph_format.space_before = Pt(vor)
    p.paragraph_format.space_after = Pt(nach)


class WordBogen:
    BREITE = 186  # nutzbare Breite in mm (A4 minus 12 mm Rand je Seite)

    def __init__(self, titel, untertitel, stufe):
        self.titel, self.untertitel = titel, untertitel
        self.gross = stufe == "gs"
        f = FARBEN[stufe]
        self.akzent, self.hell = f["akzent"].lstrip("#"), f["hell"].lstrip("#")
        self.zeile_normal = 9.5 if self.gross else 7.2
        self.zeile = self.zeile_normal
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(12)
        sec.top_margin, sec.bottom_margin = Mm(11), Mm(12)
        sec.footer_distance = Mm(6)
        st = self.doc.styles["Normal"]
        st.font.name = "Arial"
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        st.font.size = Pt(11.5 if self.gross else 9.5)
        _absatz_abstand(self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph())
        self._erster_absatz = True
        fp = sec.footer.paragraphs[0]
        r = fp.add_run(f"{titel} · {untertitel}")
        r.font.size, r.font.color.rgb = Pt(8), RGBColor(0x77, 0x77, 0x77)

    def _p(self, text="", fett=False, groesse=None, farbe=None, vor=0, nach=0):
        if self._erster_absatz:
            p = self.doc.paragraphs[0]
            self._erster_absatz = False
        else:
            p = self.doc.add_paragraph()
        _absatz_abstand(p, vor, nach)
        if text:
            r = p.add_run(text)
            r.bold = fett
            if groesse:
                r.font.size = Pt(groesse)
            if farbe:
                r.font.color.rgb = RGBColor.from_string(farbe)
        return p

    def _tabelle(self, zeilen, breiten_mm):
        t = self.doc.add_table(rows=zeilen, cols=len(breiten_mm))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        for col, b in zip(t.columns, breiten_mm):
            col.width = Mm(b)
        for row in t.rows:
            for cell, b in zip(row.cells, breiten_mm):
                cell.width = Mm(b)
                _absatz_abstand(cell.paragraphs[0])
        return t

    @staticmethod
    def _hoehe(row, mm):
        row.height = Mm(mm)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

    @staticmethod
    def _bild(run, name, hoehe_mm):
        pfad = EIGENES_LOGO if name == "logo" and EIGENES_LOGO.exists() else PNG / f"{name}.png"
        run.add_picture(str(pfad), height=Mm(hoehe_mm))

    def kopf(self, felder):
        t = self._tabelle(1, [44, 92, 50])
        logo, titel, bild = t.rows[0].cells
        for c in t.rows[0].cells:
            _rand(c)
            c.vertical_alignment = 1
        _rand(titel)
        self._bild(logo.paragraphs[0].add_run(), "logo", 21)
        p = titel.paragraphs[0]
        r = p.add_run(self.titel)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(21 if self.gross else 18), RGBColor.from_string(self.akzent)
        r = titel.add_paragraph().add_run(self.untertitel)
        r.font.size, r.font.color.rgb = Pt(10), RGBColor(0x66, 0x66, 0x66)
        bild.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._bild(bild.paragraphs[0].add_run(), "klassenrat", 23)
        for c in t.rows[0].cells:
            _rand(c, bottom=12)
        self._p().paragraph_format.line_spacing = Pt(4)
        self.felder(felder)

    def ablauf(self, schritte):
        p = self._p(vor=3, nach=2)
        r = p.add_run("Ablauf:  ")
        r.bold, r.font.size = True, Pt(9)
        r = p.add_run("   →   ".join(f"{i} {s}" for i, s in enumerate(schritte, 1)))
        r.font.size = Pt(9)

    def abschnitt(self, nr, titel, hinweis=None, icon=None):
        t = self._tabelle(1, [self.BREITE])
        c = t.rows[0].cells[0]
        _schattierung(c, self.hell)
        _rand(c, bottom=8)
        p = c.paragraphs[0]
        _absatz_abstand(p, 2, 2)
        if icon:
            self._bild(p.add_run(), icon, 7 if self.gross else 5.5)
            p.add_run("  ")
        r = p.add_run(f"{nr}  {titel}" if nr else titel)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(13 if self.gross else 11), RGBColor.from_string(self.akzent)
        if hinweis:
            r = p.add_run(f"    {hinweis}")
            r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0x55, 0x55, 0x55)
        self._p(nach=0).paragraph_format.line_spacing = Pt(4)

    def unterabschnitt(self, titel, hinweis=None, icon=None):
        p = self._p(vor=4, nach=1)
        if icon:
            self._bild(p.add_run(), icon, 5.5 if self.gross else 4.2)
            p.add_run(" ")
        r = p.add_run(titel)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(11 if self.gross else 9.5), RGBColor.from_string(self.akzent)
        if hinweis:
            r = p.add_run(f"   {hinweis}")
            r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), RGBColor(0x55, 0x55, 0x55)

    def block_anfang(self, titel, icon=None):
        t = self._tabelle(1, [self.BREITE])
        c = t.rows[0].cells[0]
        _schattierung(c, self.akzent)
        p = c.paragraphs[0]
        _absatz_abstand(p, 2, 2)
        if icon:
            self._bild(p.add_run(), icon, 6.5 if self.gross else 5)
            p.add_run("  ")
        self.zeile = 7.6 if self.gross else 5.6  # kompakter, damit zwei Blöcke auf eine Seite passen
        r = p.add_run(f"{titel}      Nr. ______")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(13 if self.gross else 11.5), RGBColor(0xFF, 0xFF, 0xFF)
        self._p().paragraph_format.line_spacing = Pt(4)

    def block_ende(self):
        self.zeile = self.zeile_normal
        self._p().paragraph_format.line_spacing = Pt(6)

    def felder(self, felder):
        gesamt = sum(f[1] for f in felder)
        breiten = []
        for label, w, *_ in felder:
            extra = 8 if len(_) and _[0] else 0
            lab = min(len(label) * (2.5 if self.gross else 2.0) + 6 + extra, self.BREITE * w / gesamt - 8)
            breiten += [lab, self.BREITE * w / gesamt - lab]
        t = self._tabelle(1, breiten)
        self._hoehe(t.rows[0], self.zeile + 1)
        for i, (label, _w, *rest) in enumerate(felder):
            lc, vc = t.rows[0].cells[2 * i], t.rows[0].cells[2 * i + 1]
            _rand(lc)
            _rand(vc, bottom=6)
            if rest and rest[0]:
                self._bild(lc.paragraphs[0].add_run(), rest[0], 6 if self.gross else 4.5)
                lc.paragraphs[0].add_run(" ")
            r = lc.paragraphs[0].add_run(label + ":")
            r.font.size = Pt(9 if not self.gross else 10.5)
            lc.vertical_alignment = vc.vertical_alignment = 3  # unten

    def tabelle(self, spalten, zeilen):
        gesamt = sum(w for _, w, _ in spalten)
        t = self._tabelle(zeilen + 1, [self.BREITE * w / gesamt for _, w, _ in spalten])
        for i, (label, _w, art) in enumerate(spalten):
            c = t.rows[0].cells[i]
            _schattierung(c, "F2F2F2")
            text = label
            if label.startswith("@"):
                icon, text = label[1:].split("|")
                self._bild(c.paragraphs[0].add_run(), icon, 5)
                text = " " + text
            r = c.paragraphs[0].add_run(text)
            r.bold, r.font.size = True, Pt(8.5 if not self.gross else 10)
            for row in t.rows[1:]:
                cell = row.cells[i]
                inhalt = self._zelle(art)
                if inhalt:
                    rr = cell.paragraphs[0].add_run(inhalt)
                    rr.font.size = Pt(8 if not self.gross else 9.5)
                    cell.vertical_alignment = 1  # Mitte
        for row in t.rows:
            for c in row.cells:
                _rand(c, top=4, bottom=4, left=4, right=4)
        for row in t.rows[1:]:
            self._hoehe(row, self.zeile)
        self._p().paragraph_format.line_spacing = Pt(3)

    @staticmethod
    def _zelle(art):
        if art.startswith("wahl:"):
            return "  ".join(f"☐ {o}" for o in art[5:].split("/"))
        if art == "gesichter":
            return "  ".join(f"☐ {t}" for _, t in GESICHTER)
        if art == "kasten":
            return "☐"
        return ""

    def linien(self, n, label=None):
        if label:
            self.unterabschnitt(label)
        t = self._tabelle(n, [self.BREITE])
        for row in t.rows:
            self._hoehe(row, self.zeile)
            _rand(row.cells[0], bottom=4)

    def ankreuzen(self, label, optionen):
        p = self._p(vor=3, nach=2)
        if label:
            r = p.add_run(label + "  ")
            r.bold = True
        p.add_run("     ".join(f"☐ {o}" for o in optionen)).font.size = Pt(9 if not self.gross else 10.5)

    def gesichter_strichliste(self, label):
        if label:
            self.unterabschnitt(label)
        t = self._tabelle(2, [self.BREITE / 3] * 3)
        for i, (_, text) in enumerate(GESICHTER):
            c = t.rows[0].cells[i]
            self._bild(c.paragraphs[0].add_run(), GESICHTER[i][0], 9)
            c.paragraphs[0].add_run("  " + text).bold = True
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for row in t.rows:
                _rand(row.cells[i], top=4, bottom=4, left=4, right=4)
        self._hoehe(t.rows[1], 16)
        self._p().paragraph_format.line_spacing = Pt(3)

    def unterschriften(self, rollen):
        self._p().paragraph_format.line_spacing = Pt(10)
        w = self.BREITE / len(rollen)
        t = self._tabelle(2, [w] * len(rollen))
        self._hoehe(t.rows[0], 12)
        for i, rolle in enumerate(rollen):
            _rand(t.rows[0].cells[i], bottom=6)
            _rand(t.rows[1].cells[i])
            r = t.rows[1].cells[i].paragraphs[0].add_run(rolle)
            r.font.size = Pt(8.5)

    def hinweis(self, text):
        p = self._p(text, groesse=8, farbe="777777", vor=6)
        p.runs[0].italic = True

    def seitenumbruch(self):
        self._p().add_run().add_break(WD_BREAK.PAGE)

    def speichern(self, pfad):
        self.doc.core_properties.title = f"{self.titel} – {self.untertitel}"
        self.doc.save(pfad)


# ---------------------------------------------------------------------------
# Inhalt der Bögen
# ---------------------------------------------------------------------------

def mittelschule(b):
    b.kopf([("Klasse", 1), ("Datum", 1.6), ("Sitzung Nr.", 1.1), ("Beginn", 1), ("Ende", 1)])
    b.ablauf(["Eröffnung", "Positivrunde", "Rückblick", "Anliegen", "Abschluss"])

    b.abschnitt("", "Rollen & Anwesenheit", icon="leute")
    b.felder([("Präsident/in", 1, "glocke"), ("Protokollant/in", 1, "stift")])
    b.felder([("Zeitwächter/in", 1, "uhr"), ("Regelwächter/in", 1, "regel")])
    b.felder([("Anwesend", 0.6), ("von", 0.4), ("Es fehlen", 2.4)])

    b.abschnitt("1", "Positivrunde", "„Mir hat gefallen, dass …“  ·  „Ich bedanke mich bei …, weil …“", icon="stern")
    b.tabelle([("Name", 1.2, ""), ("Aussage", 4.8, "")], 5)

    b.abschnitt("2", "Rückblick auf die letzte Sitzung", "Wurden unsere Beschlüsse umgesetzt?", icon="zurueck")
    b.tabelle([("Beschluss", 2.7, ""), ("zuständig", 1.2, ""),
               ("umgesetzt?", 1.9, "wahl:ja/teils/nein"), ("Bemerkung / wie geht es weiter?", 2.2, "")], 4)

    b.abschnitt("3", "Tagesordnung", "Anliegen von der Wandzeitung – Reihenfolge legt die Präsidentin / der Präsident fest",
                icon="wandzeitung")
    b.tabelle([("Nr.", 0.45, ""), ("Anliegen von", 1.5, ""), ("Thema", 3.2, ""),
               ("Art*", 0.7, ""), ("besprochen", 0.9, "kasten")], 6)
    b.hinweis("* Art:  K = Kritik  ·  W = Wunsch  ·  V = Vorschlag  ·  Ko = Konflikt  ·  I = Information.  "
              "Für jedes Anliegen wird auf den folgenden Seiten ein eigener Block ausgefüllt "
              "(Seite 2 bei Bedarf mehrfach kopieren).")
    b.seitenumbruch()

    anliegen_ms(b)
    anliegen_ms(b)
    b.seitenumbruch()

    anliegen_ms(b)
    b.abschnitt("4", "Abschluss", "Blitzlicht: Wie war die Sitzung?", icon="froh")
    b.felder([("gut", 1, "froh"), ("mittel", 1, "mittel"), ("schlecht", 1, "traurig")])
    b.felder([("Das nehmen wir uns vor", 1, "idee")])
    b.unterabschnitt("Rollen für die nächste Sitzung", icon="kalender")
    b.felder([("Präsident/in", 1, "glocke"), ("Protokollant/in", 1, "stift")])
    b.felder([("Zeitwächter/in", 1, "uhr"), ("Regelwächter/in", 1, "regel")])
    b.felder([("Nächster Klassenrat am", 1.6), ("Wandzeitung betreut", 1.4)])
    b.unterschriften(["Präsident/in", "Protokollant/in", "Lehrkraft (gesehen)"])


def anliegen_ms(b):
    b.block_anfang("Anliegen", icon="sprechen")
    b.felder([("Thema", 3), ("eingebracht von", 2)])
    b.ankreuzen("Art:", ["Kritik", "Wunsch", "Vorschlag", "Konflikt", "Information"])
    b.felder([("Worum geht es?", 1)])
    b.unterabschnitt("Aussagen / Sichtweisen", "Betroffene zuerst, dann die Klasse – sinngemäß festhalten", icon="sprechen")
    b.tabelle([("Name", 1.2, ""), ("Aussage", 4.8, "")], 5)
    b.unterabschnitt("Lösungsvorschläge & Abstimmung", icon="hand")
    b.tabelle([("Vorschlag", 3.4, ""), ("von", 1.1, ""), ("@daumen_hoch|dafür", 0.75, ""),
               ("@daumen_runter|dagegen", 0.9, ""), ("Enth.", 0.55, "")], 3)
    b.felder([("Beschluss", 1, "beschluss")])
    b.felder([("Wer?", 2), ("Bis wann?", 1.1), ("Überprüfung am", 1.2)])
    b.ankreuzen("", ["einstimmig", "Mehrheit", "vertagt", "keine Einigung",
                     "weiter an Lehrkraft / SMV / Schulleitung"])
    b.block_ende()


def grundschule(b):
    b.kopf([("Klasse", 1), ("Datum", 1.6), ("Nr.", 0.8)])
    b.ablauf(["Begrüßung", "Das war schön", "Hat es geklappt?", "Unsere Themen", "Wie war's?"])

    b.abschnitt("", "Wer macht heute was?", icon="leute")
    b.felder([("Leitung", 1, "glocke"), ("Schreiber/in", 1, "stift")])
    b.felder([("Zeitwächter/in", 1, "uhr"), ("Regelwächter/in", 1, "regel")])
    b.felder([("Es fehlen", 1)])

    b.abschnitt("1", "Das war schön", "„Ich fand schön, dass …“", icon="stern")
    b.tabelle([("Name", 1.3, ""), ("Was war schön?", 3.7, "")], 4)

    b.abschnitt("2", "Hat es geklappt?", "Das hatten wir letztes Mal beschlossen", icon="zurueck")
    b.tabelle([("Was wollten wir machen?", 3.0, ""), ("Wer?", 1.1, ""),
               ("Hat es geklappt?", 1.8, "gesichter")], 3)

    b.abschnitt("3", "Darüber sprechen wir heute", "Themen von der Wandzeitung", icon="wandzeitung")
    b.tabelle([("Nr.", 0.45, ""), ("Wer?", 1.4, ""), ("Thema", 3.6, "")], 4)
    b.hinweis("Für jedes Thema wird auf den nächsten Seiten ein eigenes Feld ausgefüllt "
              "(Seite 2 bei Bedarf mehrfach kopieren).")
    b.seitenumbruch()

    thema_gs(b)
    thema_gs(b)
    b.seitenumbruch()

    thema_gs(b)
    b.abschnitt("4", "Wie war unser Klassenrat?", "Striche machen", icon="froh")
    b.gesichter_strichliste("")
    b.unterabschnitt("Wer macht nächstes Mal was?", icon="kalender")
    b.felder([("Leitung", 1, "glocke"), ("Schreiber/in", 1, "stift")])
    b.felder([("Zeitwächter/in", 1, "uhr"), ("Regelwächter/in", 1, "regel")])
    b.unterschriften(["Leitung", "Schreiber/in", "Lehrkraft"])


def thema_gs(b):
    b.block_anfang("Thema", icon="sprechen")
    b.felder([("Thema", 3), ("Wer hat das Thema?", 2)])
    b.unterabschnitt("Was sagen die Kinder?", icon="sprechen")
    b.tabelle([("Name", 1.3, ""), ("Das sagt das Kind", 3.7, "")], 4)
    b.unterabschnitt("Unsere Ideen", icon="idee")
    b.tabelle([("Idee", 3.2, ""), ("von", 1.1, ""),
               ("@daumen_hoch|dafür", 0.85, ""), ("@daumen_runter|dagegen", 0.95, "")], 2)
    b.felder([("Das machen wir", 1, "haken")])
    b.felder([("Wer?", 2), ("Bis wann?", 1)])
    b.ankreuzen("", ["Wir sprechen nächstes Mal weiter", "Das klären wir mit der Lehrkraft"])
    b.block_ende()


BOEGEN = [
    ("protokoll-grundschule", "Unser Klassenrat", "Protokoll · Grundschule Klasse 1–4", "gs", grundschule),
    ("protokoll-mittelschule", "Klassenrat – Protokoll", "Mittelschule Klasse 5–9", "ms", mittelschule),
]


def main():
    druck, word = ROOT / "druck", ROOT / "word"
    druck.mkdir(exist_ok=True)
    word.mkdir(exist_ok=True)
    for name, titel, untertitel, stufe, inhalt in BOEGEN:
        for Klasse, pfad in ((HtmlBogen, druck / f"{name}.html"), (WordBogen, word / f"{name}.docx")):
            b = Klasse(titel, untertitel, stufe)
            inhalt(b)
            b.speichern(pfad)
            print("geschrieben:", pfad.relative_to(ROOT.parent))


if __name__ == "__main__":
    main()
